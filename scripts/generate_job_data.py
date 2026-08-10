"""Convert job-applications.xlsx into data/job-applications.json for job-tracker.html.

Usage:
    python scripts/generate_job_data.py

Reads the "Applications" and "Answers" sheets of job-applications.xlsx (in the
repo root) and writes data/job-applications.json, sorted by Applied Date
descending. Each entry gets an "answers" array nested in from the Answers
sheet, matched by Company + Role / Position.
"""

import json
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = REPO_ROOT / "job-applications.xlsx"
JSON_PATH = REPO_ROOT / "data" / "job-applications.json"
APPLICATIONS_SHEET = "Applications"
ANSWERS_SHEET = "Answers"

# Maps xlsx column headers to output JSON keys. Only these columns are used;
# any other column in the sheet (Salary, Portal Link, etc.) is ignored for now.
COLUMN_MAP = {
    "Company": "company",
    "Role / Position": "role",
    "Type": "type",
    "Location": "location",
    "Applied Date": "date_applied",
    "Deadline": "deadline",
    "Status": "status",
    "Assessment Stage": "stage",
    "Priority": "priority",
    "Reason": "reason",
    "Next Action": "next_action",
}

ANSWERS_COLUMN_MAP = {
    "Company": "company",
    "Role / Position": "role",
    "Question": "question",
    "Tags": "tags",
    "Answer Text": "text",
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if hasattr(value, "strftime"):
        # date/datetime cell (Excel stores dates as datetimes even with no
        # time component) — keep it a plain YYYY-MM-DD, not a full timestamp.
        return value.strftime("%Y-%m-%d")
    return value


def match_key(company, role):
    return (company or "").strip().lower(), (role or "").strip().lower()


def load_sheet_rows(wb, sheet_name, column_map):
    if sheet_name not in wb.sheetnames:
        sys.exit(f'Sheet "{sheet_name}" not found. Sheets present: {wb.sheetnames}')
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    col_index = {h: i for i, h in enumerate(headers) if h in column_map}

    missing = set(column_map) - set(col_index)
    if missing:
        sys.exit(f"Missing expected column(s) in \"{sheet_name}\": {sorted(missing)}")

    return rows[1:], col_index


def load_answers(wb):
    rows, col_index = load_sheet_rows(wb, ANSWERS_SHEET, ANSWERS_COLUMN_MAP)

    answers_by_key = {}
    for row in rows:
        company = clean(row[col_index["Company"]])
        role = clean(row[col_index["Role / Position"]])
        question = clean(row[col_index["Question"]])
        if not company or not role or not question:
            continue  # blank template row

        tags_raw = clean(row[col_index["Tags"]]) or ""
        tags = [t.strip() for t in tags_raw.split(",") if t.strip()]

        answer = {
            "question": question,
            "tags": tags,
            "text": clean(row[col_index["Answer Text"]]) or "",
        }
        answers_by_key.setdefault(match_key(company, role), []).append(answer)

    return answers_by_key


def main():
    if not XLSX_PATH.exists():
        sys.exit(f"Not found: {XLSX_PATH}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    answers_by_key = load_answers(wb)

    app_rows, col_index = load_sheet_rows(wb, APPLICATIONS_SHEET, COLUMN_MAP)

    entries = []
    matched_answer_keys = set()
    for row in app_rows:
        company = clean(row[col_index["Company"]])
        role = clean(row[col_index["Role / Position"]])
        if not company or not role:
            continue  # blank template row

        entry = {
            json_key: clean(row[col_index[xlsx_col]])
            for xlsx_col, json_key in COLUMN_MAP.items()
        }
        key = match_key(company, role)
        entry["answers"] = answers_by_key.get(key, [])
        matched_answer_keys.add(key)
        entries.append(entry)

    entries.sort(key=lambda e: e["date_applied"] or "", reverse=True)

    orphaned = set(answers_by_key) - matched_answer_keys
    if orphaned:
        print(
            "Warning: answers found for company/role pairs with no matching "
            f"application row (skipped): {sorted(orphaned)}"
        )

    JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    JSON_PATH.write_text(json.dumps(entries, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(entries)} entries to {JSON_PATH.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
